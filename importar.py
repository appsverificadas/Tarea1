import os
import pathlib
import openpyxl
from datetime import datetime
from sqlalchemy.orm import Session
from database import engine, Base
from models import Insumo, Movimiento, Empleado

carpeta = pathlib.Path(__file__).parent.absolute()
archivos = [f for f in os.listdir(carpeta) if f.endswith(".xlsx") and not f.startswith("~")]

if not archivos:
    print("❌ No se encontró el archivo Excel en la carpeta.")
else:
    print(f"✅ Archivo detectado: {archivos[0]} ...")
    
    try:
        documento = openpyxl.load_workbook(os.path.join(carpeta, archivos[0]), data_only=True)
        hoja = documento.worksheets[0]
        
        # Borramos la base vacía y creamos la estructura nueva completa
        Base.metadata.drop_all(bind=engine)
        Base.metadata.create_all(bind=engine)
        
        db = Session(engine)
        insumos_agregados = 0
        empleados_agregados = 0
        movimientos_agregados = 0

        # PASO 1: Importar Insumos
        for fila in hoja.iter_rows(min_row=4, values_only=True):
            if len(fila) > 17:
                nombre = fila[16]
                stock = fila[17]
                if nombre and str(nombre).strip() and str(nombre).strip() not in ('…', 'None'):
                    nombre_limpio = str(nombre).strip()
                    if not db.query(Insumo).filter(Insumo.nombre == nombre_limpio).first():
                        stock_limpio = int(stock) if stock not in (None, "") else 0
                        db.add(Insumo(nombre=nombre_limpio, stock_inicial=stock_limpio))
                        insumos_agregados += 1
        db.commit()

        # PASO 2: Extraer Empleados automáticamente del historial
        for fila in hoja.iter_rows(min_row=4, values_only=True):
            responsable = fila[5]
            if responsable and str(responsable).strip() != 'None':
                nombre_resp = str(responsable).strip().title()
                # Verifica que no lo hayamos guardado ya para no duplicar
                if not db.query(Empleado).filter(Empleado.nombre == nombre_resp).first():
                    db.add(Empleado(nombre=nombre_resp))
                    empleados_agregados += 1
        db.commit()

        # PASO 3: Importar Movimientos
        for fila in hoja.iter_rows(min_row=4, values_only=True):
            paciente = fila[1]      
            articulo = fila[2]      
            cantidad = fila[3]      
            fecha_excel = fila[4]   
            responsable = fila[5]   

            if articulo and cantidad and str(articulo).strip() != 'None':
                nombre_articulo = str(articulo).strip()
                insumo_db = db.query(Insumo).filter(Insumo.nombre == nombre_articulo).first()
                
                if insumo_db:
                    fecha_mov = fecha_excel if isinstance(fecha_excel, datetime) else datetime.now()
                    resp_limpio = str(responsable).strip().title() if responsable and str(responsable).strip() != 'None' else "-"
                    
                    nuevo_movimiento = Movimiento(
                        fecha=fecha_mov,
                        tipo="EGRESO",
                        cantidad=int(cantidad),
                        paciente=str(paciente) if paciente and str(paciente) != 'None' else "-",
                        responsable=resp_limpio
                    )
                    nuevo_movimiento.insumo_id = insumo_db.id
                    db.add(nuevo_movimiento)
                    movimientos_agregados += 1
        db.commit()
        
        print(f"🎉 ¡ÉXITO TOTAL!")
        print(f"📦 {insumos_agregados} tipos de insumos guardados.")
        print(f"🧑‍⚕️ {empleados_agregados} profesionales extraídos y guardados.")
        print(f"📉 {movimientos_agregados} consumos registrados.")
        
    except Exception as e:
        print(f"❌ Ocurrió un error: {e}")
    finally:
        db.close()